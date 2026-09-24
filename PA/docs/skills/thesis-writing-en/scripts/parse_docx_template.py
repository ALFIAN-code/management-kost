#!/usr/bin/env python3
"""Extract thesis template structure and styles from a DOCX file.

Usage:
  python parse_docx_template.py template.docx --out format-baseline.md
"""

from __future__ import annotations

import argparse
import re
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document


W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def emu_to_cm(value) -> str:
    if value is None:
        return ""
    return f"{value.cm:.2f} cm"


def font_size(value) -> str:
    if value is None:
        return ""
    return f"{value.pt:g} pt"


def has_cjk(text: str) -> bool:
    return any("\u3400" <= ch <= "\u9fff" for ch in text)


def get_rfonts(run) -> dict[str, str]:
    rpr = run._element.find(f"{W_NS}rPr")
    if rpr is None:
        return {}
    rfonts = rpr.find(f"{W_NS}rFonts")
    if rfonts is None:
        return {}
    out = {}
    for key in ("ascii", "hAnsi", "eastAsia", "hint"):
        value = rfonts.get(f"{W_NS}{key}")
        if value:
            out[key] = value
    return out


def style_name(paragraph) -> str:
    try:
        return paragraph.style.name or ""
    except Exception:
        return ""


def extract_sections(doc: Document) -> list[dict[str, str]]:
    sections = []
    for idx, sec in enumerate(doc.sections, start=1):
        sections.append(
            {
                "index": str(idx),
                "page": f"{emu_to_cm(sec.page_width)} x {emu_to_cm(sec.page_height)}",
                "margins": (
                    f"top {emu_to_cm(sec.top_margin)}, bottom {emu_to_cm(sec.bottom_margin)}, "
                    f"left {emu_to_cm(sec.left_margin)}, right {emu_to_cm(sec.right_margin)}"
                ),
                "header": emu_to_cm(sec.header_distance),
                "footer": emu_to_cm(sec.footer_distance),
            }
        )
    return sections


def infer_heading_level(text: str, style: str) -> int | None:
    style_l = style.lower()
    m = re.search(r"(?:heading|toc)\s*(\d+)", style_l)
    if m:
        return int(m.group(1))
    m = re.search(r"(\d+)\s*[-_ ]?\s*\d*级", style)
    if m:
        return int(m.group(1))
    if re.match(r"^(第[一二三四五六七八九十百]+章|\d+(\.\d+){0,3}\s+|chapter\s+\d+)", text, re.I):
        return text.count(".") + 1 if re.match(r"^\d+(\.\d+)", text) else 1
    return None


def extract_chapter_structure(doc: Document) -> list[dict[str, str]]:
    entries = []
    seen_titles = set()

    def add_entry(level: int | None, title: str, style: str, source: str) -> None:
        clean = re.sub(r"\s*\d+\s*$", "", title).strip()
        if not clean or clean in seen_titles:
            return
        seen_titles.add(clean)
        entries.append({"level": str(level or 1), "title": clean, "style": style, "source": source})

    in_toc = False
    for p in doc.paragraphs:
        text = p.text.strip()
        style = style_name(p)
        if not text:
            continue
        compact = text.replace(" ", "")
        if compact in {"目录", "目錄"} or text.lower() in {"contents", "table of contents"}:
            in_toc = True
            continue
        if in_toc:
            level = infer_heading_level(text, style)
            if "toc" in style.lower() or level is not None:
                add_entry(level, text, style, "toc")
                continue
            if len(entries) >= 2:
                break

    if entries:
        return entries

    for p in doc.paragraphs:
        text = p.text.strip()
        style = style_name(p)
        if not text:
            continue
        level = infer_heading_level(text, style)
        if level is not None and level <= 4:
            add_entry(level, text, style, "body")
    return entries


def detect_special(title: str) -> str:
    pairs = [
        ("外文资料原文|foreign material|original text", "foreign-original"),
        ("外文资料译文|translation", "foreign-translation"),
        ("附录|appendix", "appendix"),
        ("声明|授权|declaration|authorization", "declaration"),
        ("致谢|acknowledg", "acknowledgments"),
        ("参考文献|references|bibliography", "references"),
        ("摘要|abstract", "abstract"),
        ("目录|contents", "toc"),
    ]
    lower = title.lower()
    for pattern, kind in pairs:
        if re.search(pattern, lower, re.I):
            return kind
    return "body"


def extract_styles(doc: Document) -> list[dict[str, str]]:
    samples: dict[str, dict[str, str]] = {}
    counts = Counter()
    body_candidates = Counter()
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = style_name(p)
        counts[style] += 1
        if len(text) >= 80:
            body_candidates[style] += 1
        if style in samples:
            continue
        run = next((r for r in p.runs if r.text.strip()), None)
        samples[style] = {
            "style": style,
            "sample": text[:80].replace("|", "\\|"),
            "font": run.font.name if run else "",
            "size": font_size(run.font.size) if run else "",
            "bold": str(run.font.bold) if run else "",
            "rfonts": ", ".join(f"{k}={v}" for k, v in get_rfonts(run).items()) if run else "",
            "cjk_hint": "yes" if run and has_cjk(run.text) and get_rfonts(run).get("hint") == "eastAsia" else "",
            "count": str(counts[style]),
        }
    for style, count in counts.items():
        if style in samples:
            samples[style]["count"] = str(count)
            samples[style]["body_candidate"] = "yes" if body_candidates[style] else ""
    return sorted(samples.values(), key=lambda item: (-int(item["count"]), item["style"]))


def extract_inline_constraints(doc: Document) -> list[str]:
    patterns = [
        r"不少于\s*\d+\s*[篇字词]",
        r"\d+\s*[字词]\s*(左右|以内|以上|以下)",
        r"at least\s+\d+",
        r"no more than\s+\d+",
        r"\d+\s+words?",
        r"\d+\s+references?",
    ]
    found = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        for pattern in patterns:
            if re.search(pattern, text, re.I):
                found.append(text[:180])
                break
    return found[:30]


def extract_cover_candidates(doc: Document) -> list[str]:
    labels = re.compile(r"(题目|学院|专业|姓名|学号|教师|日期|title|college|major|student|supervisor|date)", re.I)
    out = []
    for p in doc.paragraphs[:80]:
        text = p.text.strip()
        if text and labels.search(text):
            out.append(text[:160])
    return out[:30]


def render_markdown(path: Path, doc: Document) -> str:
    lines = ["# Format Baseline", ""]

    lines += ["## Source", f"- Template: `{path}`", ""]

    lines += ["## Page and Sections", ""]
    lines += ["| # | Page | Margins | Header | Footer |", "|---|---|---|---|---|"]
    for sec in extract_sections(doc):
        lines.append(f"| {sec['index']} | {sec['page']} | {sec['margins']} | {sec['header']} | {sec['footer']} |")
    lines.append("")

    lines += ["## Chapter Structure", ""]
    lines += ["| Level | Title | Type | Style | Source |", "|---|---|---|---|---|"]
    for item in extract_chapter_structure(doc):
        kind = detect_special(item["title"])
        lines.append(f"| H{item['level']} | {item['title']} | {kind} | {item['style']} | {item['source']} |")
    lines.append("")

    lines += ["## Style Inventory", ""]
    lines += ["| Style | Count | Body Candidate | Font | Size | Bold | Run Fonts | CJK Hint | Sample |"]
    lines += ["|---|---:|---|---|---|---|---|---|---|"]
    for item in extract_styles(doc):
        lines.append(
            f"| {item['style']} | {item['count']} | {item.get('body_candidate', '')} | "
            f"{item['font']} | {item['size']} | {item['bold']} | {item['rfonts']} | "
            f"{item['cjk_hint']} | {item['sample']} |"
        )
    lines.append("")

    constraints = extract_inline_constraints(doc)
    lines += ["## Inline Constraints", ""]
    if constraints:
        lines += [f"- {item}" for item in constraints]
    else:
        lines.append("- Not detected.")
    lines.append("")

    cover = extract_cover_candidates(doc)
    lines += ["## Cover Field Candidates", ""]
    if cover:
        lines += [f"- {item}" for item in cover]
    else:
        lines.append("- Not detected.")
    lines.append("")

    lines += [
        "## Manual Confirmation Needed",
        "- Confirm the semantic style mapping for level-1/2/3 headings, body text, captions, references, and TOC.",
        "- Confirm citation style and minimum reference count if not visible above.",
        "- Confirm whether special sections are required and whether they should be copied from the template.",
        "",
    ]
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("template", type=Path)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()

    doc = Document(args.template)
    md = render_markdown(args.template, doc)
    if args.out:
        args.out.write_text(md, encoding="utf-8")
    else:
        print(md)


if __name__ == "__main__":
    main()
