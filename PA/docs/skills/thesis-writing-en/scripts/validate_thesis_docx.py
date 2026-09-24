#!/usr/bin/env python3
"""Validate final thesis DOCX details that commonly regress in WPS/Word.

Checks include numeric citation order, citation superscript, keyword counts,
reference auto-numbering, TOC field contents, and heading outline leakage.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

CITATION_RE = re.compile(r"\[([0-9]{1,3})\]")


def body_paragraphs_until_references(doc: Document, ref_heading: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == ref_heading:
            break
        yield paragraph


def first_seen_citations(doc: Document, ref_heading: str) -> tuple[list[int], tuple[int, str] | None]:
    seen: set[int] = set()
    order: list[int] = []
    first: tuple[int, str] | None = None
    for index, paragraph in enumerate(body_paragraphs_until_references(doc, ref_heading), 1):
        for match in CITATION_RE.finditer(paragraph.text):
            if first is None:
                first = (index, paragraph.text[:220])
            number = int(match.group(1))
            if number not in seen:
                seen.add(number)
                order.append(number)
    return order, first


def citation_superscript_stats(doc: Document, ref_heading: str):
    total = 0
    superscript = 0
    bad: list[tuple[int, str, str]] = []
    for index, paragraph in enumerate(body_paragraphs_until_references(doc, ref_heading), 1):
        for run in paragraph.runs:
            if CITATION_RE.search(run.text):
                total += 1
                rpr = run._element.rPr
                vert = rpr.find(qn("w:vertAlign")) if rpr is not None else None
                if vert is not None and vert.get(qn("w:val")) == "superscript":
                    superscript += 1
                else:
                    bad.append((index, run.text, paragraph.text[:120]))
    return total, superscript, bad


def reference_paragraphs(doc: Document, ref_heading: str):
    started = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == ref_heading:
            started = True
            continue
        if not started:
            continue
        if text:
            yield paragraph


def numbering_info(paragraph):
    ppr = paragraph._element.pPr
    num_pr = ppr.find(qn("w:numPr")) if ppr is not None else None
    if num_pr is None:
        return None, None
    num_id_el = num_pr.find(qn("w:numId"))
    ilvl_el = num_pr.find(qn("w:ilvl"))
    num_id = num_id_el.get(qn("w:val")) if num_id_el is not None else None
    ilvl = ilvl_el.get(qn("w:val")) if ilvl_el is not None else None
    return num_id, ilvl


def toc_instructions(doc: Document) -> list[str]:
    values: list[str] = []
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for node in run._element.iter():
                if node.tag == qn("w:instrText") and node.text:
                    values.append(node.text)
    return values


def style_has_outline(doc: Document, style_name: str) -> bool:
    if style_name not in [style.name for style in doc.styles]:
        return False
    ppr = doc.styles[style_name]._element.pPr
    return ppr.find(qn("w:outlineLvl")) is not None if ppr is not None else False


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    parser.add_argument("--ref-heading", default="References")
    parser.add_argument("--max-cn-keywords", type=int, default=None)
    parser.add_argument("--forbid-en-keywords", action="store_true")
    parser.add_argument("--refs-auto-numbered", action="store_true")
    parser.add_argument("--toc-must-contain", action="append", default=[])
    parser.add_argument("--toc-must-not-contain", action="append", default=[])
    parser.add_argument("--no-outline-style", action="append", default=[])
    args = parser.parse_args()

    doc = Document(str(args.docx))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    failures: list[str] = []

    print(f"DOCX: {args.docx}")
    print(f"Paragraphs: {len(doc.paragraphs)}  Tables: {len(doc.tables)}  Sections: {len(doc.sections)}")

    cn_keywords = full_text.count("关键词：")
    en_keywords = full_text.count("Keywords:")
    print(f"Keywords: Chinese={cn_keywords} English={en_keywords}")
    if args.max_cn_keywords is not None and cn_keywords > args.max_cn_keywords:
        failures.append(f"Chinese keyword lines exceed {args.max_cn_keywords}: {cn_keywords}")
    if args.forbid_en_keywords and en_keywords:
        failures.append(f"English Keywords lines are forbidden but found: {en_keywords}")

    order, first = first_seen_citations(doc, args.ref_heading)
    expected = list(range(1, len(order) + 1))
    print(f"First citation: {first}")
    print(f"First-seen citation order: {order}")
    if order != expected:
        failures.append(f"Citation first-seen order must be {expected}, got {order}")

    total, superscript, bad = citation_superscript_stats(doc, args.ref_heading)
    print(f"Citation superscript: {superscript}/{total}")
    if total != superscript:
        failures.append(f"Some body citation runs are not superscript: {bad[:3]}")

    refs = list(reference_paragraphs(doc, args.ref_heading))
    if refs:
        first_ref = refs[0]
        num_id, ilvl = numbering_info(first_ref)
        literal_prefixes = [p.text[:30] for p in refs if re.match(r"^\s*\[\d+\]", p.text)]
        print(f"First reference style={first_ref.style.name} numId={num_id} ilvl={ilvl} literal_prefix={bool(literal_prefixes)}")
        if args.refs_auto_numbered:
            if literal_prefixes:
                failures.append(f"Reference paragraphs contain manual numeric prefixes: {literal_prefixes[:3]}")
            if num_id is None:
                failures.append("Reference paragraphs do not appear to use Word numbering")
    else:
        failures.append(f"No reference paragraphs found after heading {args.ref_heading!r}")

    instr = toc_instructions(doc)
    toc = [item for item in instr if "TOC" in item]
    print(f"TOC instructions: {toc}")
    toc_joined = "\n".join(toc)
    for token in args.toc_must_contain:
        if token not in toc_joined:
            failures.append(f"TOC instruction does not contain required token: {token}")
    for token in args.toc_must_not_contain:
        if token in toc_joined:
            failures.append(f"TOC instruction contains forbidden token: {token}")

    for style_name in args.no_outline_style:
        has_outline = style_has_outline(doc, style_name)
        print(f"Style {style_name!r} has outline level: {has_outline}")
        if has_outline:
            failures.append(f"Style {style_name!r} still has outline level")

    if failures:
        print("\nFAIL", file=sys.stderr)
        for failure in failures:
            print(f"- {failure}", file=sys.stderr)
        return 1

    print("OK thesis DOCX validation passed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
