#!/usr/bin/env python3
"""Extract reviewable Markdown from a DOCX thesis draft.

Usage:
  python extract_docx_draft.py draft.docx --out draft-extracted.md
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document


def is_heading(text: str, style: str) -> bool:
    style_l = style.lower()
    if any(key in style_l for key in ("heading", "title", "标题")):
        return True
    return bool(re.match(r"^(第[一二三四五六七八九十百]+章|\d+(\.\d+){0,3}\s+|chapter\s+\d+)", text, re.I))


def heading_level(text: str, style: str) -> int:
    m = re.search(r"heading\s*(\d+)", style.lower())
    if m:
        return min(int(m.group(1)), 6)
    if re.match(r"^\d+(\.\d+)+", text):
        return min(text.split()[0].count(".") + 1, 6)
    if re.match(r"^(第[一二三四五六七八九十百]+章|chapter\s+\d+)", text, re.I):
        return 1
    return 2


def table_to_markdown(table) -> list[str]:
    rows = []
    for row in table.rows:
        rows.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
    if not rows:
        return []
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    for row in rows[1:]:
        out.append("| " + " | ".join(row) + " |")
    return out


def iter_block_items(doc: Document):
    body = doc.element.body
    paragraphs = {p._p: p for p in doc.paragraphs}
    tables = {t._tbl: t for t in doc.tables}
    for child in body.iterchildren():
        if child in paragraphs:
            yield "paragraph", paragraphs[child]
        elif child in tables:
            yield "table", tables[child]


def extract(path: Path) -> str:
    doc = Document(path)
    lines = [f"# Extracted Draft", "", f"- Source: `{path}`", ""]
    paragraph_count = 0
    table_count = 0

    for kind, block in iter_block_items(doc):
        if kind == "paragraph":
            text = block.text.strip()
            if not text:
                continue
            paragraph_count += 1
            style = block.style.name if block.style else ""
            if is_heading(text, style):
                level = heading_level(text, style)
                lines += ["", "#" * level + " " + text, f"<!-- style: {style} -->", ""]
            else:
                lines += [text, f"<!-- style: {style} -->", ""]
        else:
            table_count += 1
            lines += ["", f"<!-- table {table_count} -->"]
            lines += table_to_markdown(block)
            lines.append("")

    lines += [
        "",
        "## Extraction Summary",
        f"- Paragraphs: {paragraph_count}",
        f"- Tables: {table_count}",
    ]
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("draft", type=Path)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()

    md = extract(args.draft)
    if args.out:
        args.out.write_text(md, encoding="utf-8")
    else:
        print(md)


if __name__ == "__main__":
    main()

